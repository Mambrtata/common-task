#!/usr/bin/env python3
"""Krok 3: vyhodnotenie výťažnosti – koľko stiahnutých ponúk obsahuje
strojovo čitateľný výkaz výmer.

Použitie:
    python3 03_vyhodnotenie.py

Prejde data/subory/, pri ZIP-och nazrie dovnútra, pri XLSX súboroch hľadá
hárky s položkami výkazu výmer (TSKP kód + MJ + množstvo + cena).
Vyžaduje: pip install openpyxl
Výstup: data/report.md + ukážka položiek data/ukazka_poloziek.csv
"""

import csv
import io
import pathlib
import re
import zipfile
from collections import Counter

try:
    import openpyxl
except ImportError:
    raise SystemExit("chýba openpyxl: pip install openpyxl")

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None   # PDF analýza sa preskočí

HERE = pathlib.Path(__file__).parent
DATA = HERE / "data"

# TSKP kód položky: 6–9 číslic, príp. s bodkami/pomlčkami (napr. 311238114)
RE_TSKP = re.compile(r"^\d{6,9}([.-]\w+)?$")
MJ_ZNAME = {"m", "m2", "m3", "kus", "ks", "t", "kg", "sub", "súb", "bm", "hod",
            "m^2", "m^3", "mj"}


def najdi_xlsx(cesta: pathlib.Path):
    """Vráti zoznam (meno, bytes) xlsx súborov – priamo alebo vnútri ZIPu."""
    pripona = cesta.suffix.lower()
    if pripona in (".xlsx", ".xlsm"):
        return [(cesta.name, cesta.read_bytes())]
    if pripona == ".zip":
        vysledok = []
        try:
            with zipfile.ZipFile(cesta) as z:
                for n in z.namelist():
                    if n.lower().endswith((".xlsx", ".xlsm")):
                        vysledok.append((n, z.read(n)))
        except zipfile.BadZipFile:
            pass
        return vysledok
    return []


def extrahuj_polozky(data: bytes):
    """Z xlsx vytiahni riadky, ktoré vyzerajú ako položky výkazu výmer.

    Heuristika: v riadku je TSKP-like kód, známa MJ a aspoň 2 čísla
    (množstvo, jednotková cena). Vracia list dictov.
    """
    polozky = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                    data_only=True)
    except Exception:
        return polozky
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            hodnoty = [c for c in row if c is not None]
            if len(hodnoty) < 4:
                continue
            kod = next((str(c).strip() for c in hodnoty
                        if RE_TSKP.match(str(c).strip())), None)
            if not kod:
                continue
            mj = next((str(c).strip().lower() for c in hodnoty
                       if str(c).strip().lower() in MJ_ZNAME), None)
            cisla = [c for c in hodnoty
                     if isinstance(c, (int, float)) and not isinstance(c, bool)]
            popis = next((str(c) for c in hodnoty
                          if isinstance(c, str) and len(c) > 12
                          and not RE_TSKP.match(c.strip())), "")
            if mj and len(cisla) >= 2:
                polozky.append({"harok": ws.title, "kod": kod, "popis": popis[:90],
                                "mj": mj, "cisla": cisla[:4]})
    wb.close()
    return polozky


# riadok položky v PDF rozpočte (export z KROS): kód, popis, MJ, čísla
RE_PDF_POLOZKA = re.compile(
    r"(\d{6,9}(?:\.\w+)?)\s+(.{10,120}?)\s+"
    r"(m2|m3|m|ks|kus|t|kg|súb|sub|bm|hod)\s+"
    r"([\d\s]+(?:[.,]\d+)?)\s+([\d\s]+(?:[.,]\d+)?)", re.I)


def extrahuj_polozky_pdf(cesta: pathlib.Path):
    """Z textového PDF rozpočtu vytiahni položky. Vráti (polozky, ma_text)."""
    polozky, ma_text = [], False
    if pdfium is None:
        return polozky, ma_text
    try:
        pdf = pdfium.PdfDocument(str(cesta))
    except Exception:
        return polozky, ma_text
    try:
        for i in range(len(pdf)):
            try:
                text = pdf[i].get_textpage().get_text_range()
            except Exception:
                continue    # poškodená stránka/PDF – preskoč
            if len(text) > 100:
                ma_text = True
            for m in RE_PDF_POLOZKA.finditer(text.replace("\r", "\n")):
                polozky.append({"harok": "pdf", "kod": m.group(1),
                                "popis": m.group(2).strip()[:90],
                                "mj": m.group(3).lower(),
                                "cisla": [m.group(4).replace(" ", ""),
                                          m.group(5).replace(" ", "")]})
    finally:
        pdf.close()
    return polozky, ma_text


def main():
    subory_dir = DATA / "subory"
    if not subory_dir.exists():
        raise SystemExit("najprv spusti 02_dokumenty.py")

    stats = Counter()
    zakazky_stav = {}
    ukazka = []

    for zdir in sorted(subory_dir.iterdir()):
        if not zdir.is_dir():
            continue
        stav = {"suborov": 0, "xlsx": 0, "vykazov": 0, "poloziek": 0,
                "pdf_text": 0, "pdf_sken": 0}
        for f in zdir.rglob("*"):
            if not f.is_file():
                continue
            stav["suborov"] += 1
            stats["suborov"] += 1
            stats[f"pripona:{f.suffix.lower() or 'bez'}"] += 1
            for meno, data in najdi_xlsx(f):
                stav["xlsx"] += 1
                polozky = extrahuj_polozky(data)
                if len(polozky) >= 10:      # aspoň 10 položiek = reálny výkaz
                    stav["vykazov"] += 1
                    stav["poloziek"] += len(polozky)
                    for p in polozky[:5]:
                        ukazka.append({"zakazka_id": zdir.name, "subor": meno,
                                       **{k: str(v) for k, v in p.items()}})
            if f.suffix.lower() == ".pdf":
                polozky, ma_text = extrahuj_polozky_pdf(f)
                stav["pdf_text" if ma_text else "pdf_sken"] += 1
                if ma_text:
                    stats["pdf_textovych"] += 1
                if len(polozky) >= 10:
                    stav["vykazov"] += 1
                    stav["poloziek"] += len(polozky)
                    for p in polozky[:5]:
                        ukazka.append({"zakazka_id": zdir.name, "subor": f.name,
                                       **{k: str(v) for k, v in p.items()}})
        zakazky_stav[zdir.name] = stav

    uspesne = sum(1 for s in zakazky_stav.values() if s["vykazov"] > 0)
    celkom = len(zakazky_stav)

    r = ["# Pilot – výťažnosť extrakcie výkazov výmer", ""]
    r.append(f"- zákaziek so stiahnutými súbormi: **{celkom}**")
    r.append(f"- zákaziek s aspoň 1 čitateľným výkazom výmer: **{uspesne}** "
             f"({100 * uspesne // max(celkom, 1)} %)")
    r.append(f"- súborov spolu: {stats['suborov']}")
    r.append(f"- PDF s textovou vrstvou: {stats['pdf_textovych']} "
             f"(zvyšok PDF sú skeny -> OCR)")
    r.append(f"- položiek extrahovaných spolu (xlsx + textové PDF): "
             f"{sum(s['poloziek'] for s in zakazky_stav.values())}")
    r.append("")
    r.append("## Prípony súborov")
    for k, v in stats.most_common():
        if k.startswith("pripona:"):
            r.append(f"- {k.split(':', 1)[1]}: {v}")
    r.append("")
    r.append("## Po zákazkách")
    r.append("")
    r.append("| zákazka | súborov | xlsx | pdf text | pdf sken | výkazov | položiek |")
    r.append("|---|---|---|---|---|---|---|")
    for zid, s in zakazky_stav.items():
        r.append(f"| {zid} | {s['suborov']} | {s['xlsx']} | {s['pdf_text']} "
                 f"| {s['pdf_sken']} | {s['vykazov']} | {s['poloziek']} |")

    (DATA / "report.md").write_text("\n".join(r), encoding="utf-8")
    with (DATA / "ukazka_poloziek.csv").open("w", newline="",
                                             encoding="utf-8") as f:
        if ukazka:
            w = csv.DictWriter(f, fieldnames=list(ukazka[0]))
            w.writeheader()
            w.writerows(ukazka)
    print("\n".join(r))
    print(f"\nreport -> {DATA / 'report.md'}")


if __name__ == "__main__":
    main()
